"""Shared helpers for turning resolved rows into the upload files."""

from __future__ import annotations

import io
from copy import copy
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from ..template_store import get_template_bytes


def today_stamp() -> str:
    return datetime.now().strftime("%Y%m%d")


def parse_iso_date(text: str) -> date:
    try:
        return datetime.strptime(text.strip(), "%Y-%m-%d").date()
    except (ValueError, AttributeError) as exc:
        raise ValueError(f"'{text}' is not a valid date (expected YYYY-MM-DD).") from exc


def ms4_date(value: date) -> str:
    """MS4 files want the date as YYYYMMDD."""
    return value.strftime("%Y%m%d")


def ccs_vendor(ms4_vendor_code: str) -> str:
    """CCS files carry the MS4 vendor code with one extra leading zero."""
    return "0" + (ms4_vendor_code or "").strip()


# --------------------------------------------------------------------------
# xlsx templates
# --------------------------------------------------------------------------


class TemplateSheet:
    """A template workbook with its sample rows removed.

    The header row keeps every bit of its original formatting, and the style of
    the template's first data row is reused for each generated row so fonts,
    alignment and number formats stay exactly as the upload system expects.
    """

    def __init__(self, template_key: str, column_count: int):
        data = get_template_bytes(template_key)
        self.wb = openpyxl.load_workbook(io.BytesIO(data))
        self.ws = self.wb[self.wb.sheetnames[0]]
        self.column_count = column_count

        self._proto_style = []
        self._proto_format = []
        for col in range(1, column_count + 1):
            cell = self.ws.cell(row=2, column=col)
            self._proto_style.append(copy(cell._style))
            self._proto_format.append(cell.number_format)

        if self.ws.max_row > 1:
            self.ws.delete_rows(2, self.ws.max_row - 1)
        self._next_row = 2

    def append(self, values: list) -> None:
        row = self._next_row
        for idx, value in enumerate(values):
            cell = self.ws.cell(row=row, column=idx + 1)
            cell.value = value
            cell._style = copy(self._proto_style[idx])
            cell.number_format = self._proto_format[idx]
        self._next_row += 1

    @property
    def row_count(self) -> int:
        return self._next_row - 2

    def to_bytes(self) -> bytes:
        last = max(self._next_row - 1, 1)
        if self.ws.auto_filter.ref:
            self.ws.auto_filter.ref = (
                f"A1:{get_column_letter(self.column_count)}{last}"
            )
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


# --------------------------------------------------------------------------
# delimited templates (MS4)
# --------------------------------------------------------------------------


def template_header(template_key: str, delimiter: str) -> list[str]:
    """The header line of a csv/txt template, split into field names."""
    raw = get_template_bytes(template_key).decode("utf-8-sig", errors="replace")
    first = raw.replace("\r\n", "\n").split("\n", 1)[0]
    return first.split(delimiter)


def build_delimited(header: list[str], rows: list[list[str]], delimiter: str) -> bytes:
    """Rebuild a template-shaped file: original header, CRLF endings, no BOM."""
    lines = [delimiter.join(header)]
    for row in rows:
        lines.append(delimiter.join("" if v is None else str(v) for v in row))
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")
